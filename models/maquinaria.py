import base64
from odoo import fields, models, api
from odoo.exceptions import ValidationError


class TallerMaquinaria(models.Model):
    _name = 'taller.maquinaria'
    _description = 'Gestion maquinas'
    _rec_name = 'num_serie'
    _sql_constraints = [
        ('puerto_uniq', 'unique(puerto)', 'El puerto debe ser unico.'),
    ]

    cantidad = fields.Integer(
        string='Cantidad',
        readonly=True,
        copy=False,
        help='Numero sucesivo de la maquina'
    )
    puerto = fields.Char(
        string='Puerto',
        required=True,
        size=5,
        help='Puerto de 5 digitos. Debe ser unico.'
    )
    num_serie = fields.Char(
        string='N Serie',
        required=True
    )
    modelo_maquina = fields.Selection([
        ('consola_c10', 'Consola C10'),
        ('validadora_v5', 'Validadora V5'),
        ('validadora_v7', 'Validadora V7'),
        ('velite', 'VéLite'),
        ('ecolite_emv', 'EcoLite EMV'),
        ('router_elite', 'Router éLite'),
    ], string='Modelo', required=True, default='validadora_v5')

    cliente_taller_id = fields.Many2one(
        'taller.cliente',
        string='Cliente',
        required=True,
        ondelete='cascade'
    )
    cliente_id = fields.Many2one(
        'res.partner',
        string='Compania',
        related='cliente_taller_id.cliente_id',
        readonly=True,
        store=True
    )
    ns_cliente = fields.Char(
        string='N/S Cliente',
        compute='_compute_ns_cliente',
        store=True,
        readonly=True,
        help='Formato: Cliente-Modelo-Numero (ej: Alvarez-V5-0001)'
    )

    fecha_envio = fields.Date(
        string='F. Envio',
        help='Fecha de envio'
    )

    observaciones = fields.Text(
        string='Observaciones'
    )
    word_file = fields.Binary(
        string='Archivo Word',
        readonly=True,
        attachment=True
    )
    word_filename = fields.Char(
        string='Nombre del archivo Word',
        readonly=True
    )

    conjunto = fields.Char(
        string='Conjunto',
        help='Conjunto al que pertenece (si procede)'
    )
    raspberry = fields.Char(
        string='Raspberry',
        help='Numero de Raspberry'
    )
    id_feig = fields.Char(
        string='ID FEIG',
        help='ID Feig'
    )
    version_pcb = fields.Char(
        string='Version PCB',
        help='Version de la PCB'
    )

    estado = fields.Selection([
        ('fabricacion', 'En fabricacion'),
        ('terminada', 'Terminada'),
        ('reparacion', 'Reparacion')
    ], string='Estado', required=True, default='fabricacion')

    responsable = fields.Many2one(
        'res.users',
        string='Responsable',
        domain="[('groups_id.name', 'in', ['Taller'])]"
    )

    @api.constrains('puerto')
    def _check_puerto(self):
        for record in self:
            if record.puerto:
                if not record.puerto.isdigit():
                    raise ValidationError(
                        'El puerto debe contener solo numeros.'
                    )
                if len(record.puerto) != 5:
                    raise ValidationError(
                        'El puerto debe tener exactamente 5 digitos.'
                    )

    @api.depends('cliente_taller_id', 'modelo_maquina', 'num_serie')
    def _compute_ns_cliente(self):
        for record in self:
            cliente = record.cliente_taller_id.name or 'SinCliente'

            codigos = {
                'consola_c10': 'C10',
                'validadora_v5': 'V5',
                'validadora_v7': 'V7',
                'velite': 'VL',
                'ecolite_emv': 'EL',
                'router_elite': 'RE',
            }
            modelo = codigos.get(record.modelo_maquina, 'MAQ')

            numero = ''
            if record.num_serie:
                numero = record.num_serie[-4:] if len(record.num_serie) >= 4 else record.num_serie

            record.ns_cliente = f"{cliente}-{modelo}-{numero}"

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if not vals.get('cantidad'):
                ultimo = self.search([], order='cantidad desc', limit=1)
                vals['cantidad'] = (ultimo.cantidad or 0) + 1
        return super().create(vals_list)

    def _generate_word(self):
        """Genera un archivo Word con el contenido de observaciones."""
        import io
        from docx import Document

        for record in self:
            if not record.observaciones:
                record.word_file = False
                record.word_filename = False
                continue

            # Crear documento Word
            doc = Document()
            doc.add_heading(f'Observaciones - {record.num_serie or "Sin serie"}', level=1)
            doc.add_paragraph(f'Cliente: {record.cliente_taller_id.name or "Sin cliente"}')
            doc.add_paragraph(f'Modelo: {dict(record._fields["modelo_maquina"].selection).get(record.modelo_maquina, "")}')
            doc.add_paragraph(f'Puerto: {record.puerto or ""}')
            doc.add_paragraph(f'N Serie: {record.num_serie or ""}')
            doc.add_paragraph('')
            doc.add_heading('Observaciones:', level=2)
            doc.add_paragraph(record.observaciones)

            # Guardar en memoria
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            # Asignar al campo Binary
            record.word_file = base64.b64encode(buffer.read())
            record.word_filename = f'observaciones_{record.num_serie or record.id}.docx'

    def action_descargar_word(self):
        self.ensure_one()
        if not self.word_file:
            raise ValidationError("No hay archivo Word generado. No hay observaciones")
        return {
            'type': 'ir.actions.act_url',
            'url': (
                f'/web/content/?model=taller.maquinaria'
                f'&id={self.id}'
                f'&field=word_file'
                f'&filename_field=word_filename'
                f'&download=true'
            ),
            'target': 'self',
        }


    def write(self, vals):
        res = super().write(vals)
        if 'observaciones' in vals:
            self._generate_word()
        return res

    @api.constrains('estado', 'responsable')
    def _check_responsable_reparacion(self):
        for record in self:
            if record.estado == 'reparacion' and not record.responsable:
                raise ValidationError(
                    'Debe seleccionar un Responsable cuando el estado es Reparacion.'
                )